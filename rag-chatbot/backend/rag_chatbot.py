"""
Clario Backend API (Production-Ready).

Provides document ingestion, local vector search with ChromaDB,
hybrid lexical/semantic retrieval, and streaming grounded chat responses.
"""

import asyncio
import csv
import hashlib
import ipaddress
import json
import logging
import math
import os
import re
import socket
import tempfile
import uuid
from collections import Counter
from typing import AsyncGenerator, Dict, List, Optional
from urllib.parse import urlparse

import aiofiles
import chromadb
import docx
import httpx
import pdfplumber
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from fastapi import APIRouter, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel, Field

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

from gemini_utils import (
    GeminiConfigurationError,
    gemini_chat,
    gemini_chat_stream,
    gemini_embed,
    is_gemini_configured,
)

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", 800))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", 150))
TOP_K_RESULTS = int(os.getenv("TOP_K_RESULTS", 5))
LOCAL_EMBED_DIM = int(os.getenv("LOCAL_EMBED_DIM", 768))
LLM_MODEL = os.getenv("LLM_MODEL", "gemini-2.5-flash")
MAX_FILE_SIZE_BYTES = int(os.getenv("MAX_FILE_SIZE_BYTES", 31457280))  # 30 MB default

VECTOR_STORE_PATH_RAW = os.getenv("VECTOR_STORE_PATH", "./vector_store")
VECTOR_STORE_PATH = (
    VECTOR_STORE_PATH_RAW
    if os.path.isabs(VECTOR_STORE_PATH_RAW)
    else os.path.join(BASE_DIR, VECTOR_STORE_PATH_RAW)
)
COLLECTION_NAME = "documents"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)s  %(name)s  |  %(message)s",
)
log = logging.getLogger("clario")

app = FastAPI(
    title="Clario API",
    version="2.1.0",
    description="High-performance document ingestion and grounded chat API with Gemini & ChromaDB",
)

# CORS configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

router = APIRouter()


# ---------------------------------------------------------------------------
# Data Models
# ---------------------------------------------------------------------------

class ChatRequest(BaseModel):
    question: str
    chat_history: List[dict] = Field(default_factory=list)


class ChatResponse(BaseModel):
    answer: str
    sources: List[dict] = Field(default_factory=list)
    chunks: List[dict] = Field(default_factory=list)


class IngestResponse(BaseModel):
    status: str = "success"
    chunks_stored: int = 0
    filename: str = ""
    filenames: List[str] = Field(default_factory=list)
    files_processed: int = 0


class IngestURLRequest(BaseModel):
    url: str


class DocumentSummary(BaseModel):
    source: str
    chunks: int


class DocumentsResponse(BaseModel):
    documents: List[DocumentSummary] = Field(default_factory=list)
    total_documents: int = 0
    total_chunks: int = 0


class DeleteDocumentRequest(BaseModel):
    source: str


class DeleteDocumentResponse(BaseModel):
    status: str = "success"
    source: str
    chunks_deleted: int = 0
    documents_remaining: int = 0


# ---------------------------------------------------------------------------
# Prompts & Splitting
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """You are Clario, an expert, warm, and precise document assistant.

CRITICAL RULES:
1. Answer ONLY using the provided CONTEXT below.
2. If the answer is NOT present or cannot be inferred from the context, respond clearly:
   "I don't have enough information to answer that from the provided documents."
3. Never invent facts, hallucinate citations, or make assumptions outside the provided context.
4. Clean formatting: Do NOT append the document filename (e.g. "(filename.pdf)") to every sentence, line, or bullet point. The UI already tracks and cites all source documents automatically in a dedicated drawer below the response. Only mention a document name if contrasting multiple conflicting sources or if explicitly requested.
5. Provide a well-structured, natural explanation using Markdown headings, bullet points, or code formatting when helpful.
"""


def build_prompt(
    retrieved_chunks: List[Dict],
    chat_history: List[dict],
    question: str,
) -> str:
    history_lines = []
    for message in chat_history[-6:]:
        role = str(message.get("role", "user")).upper()
        content = str(message.get("content", "")).strip()
        if content:
            history_lines.append(f"{role}: {content}")

    history_str = "\n".join(history_lines)
    chunks_str = "\n\n---\n\n".join(
        (
            f"[Source: {chunk.get('source', 'unknown')} | Chunk #{chunk.get('chunk_index', 0)}]\n"
            f"{chunk.get('text', '')}"
        )
        for chunk in retrieved_chunks
    )
    return f"""{SYSTEM_PROMPT}

CONTEXT:
{chunks_str}

CHAT HISTORY:
{history_str or "(no previous messages)"}

QUESTION:
{question}

ANSWER:"""


def split_text(text: str, source_file: str = "unknown") -> List[Dict]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        length_function=lambda s: len(s.split()),
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    raw_chunks = splitter.split_text(text)
    return [
        {
            "text": chunk,
            "meta": {
                "source": source_file,
                "chunk_index": i,
                "id": str(uuid.uuid4()),
            },
        }
        for i, chunk in enumerate(raw_chunks)
    ]


# ---------------------------------------------------------------------------
# Vector Store & Embedding Management
# ---------------------------------------------------------------------------

os.makedirs(VECTOR_STORE_PATH, exist_ok=True)
chroma_client = chromadb.PersistentClient(path=VECTOR_STORE_PATH)


def _get_collection():
    return chroma_client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )


def _embedding_dimension() -> int:
    collection = _get_collection()
    if collection.count() == 0:
        return LOCAL_EMBED_DIM

    try:
        result = collection.get(limit=1, include=["embeddings"])
        embeddings = result.get("embeddings")
        if embeddings is not None and len(embeddings) > 0:
            return len(embeddings[0])
    except Exception:
        log.warning("Could not read collection embedding dimension; using %d", LOCAL_EMBED_DIM)
    return LOCAL_EMBED_DIM


def _tokens(text: str) -> List[str]:
    normalized = text.casefold()
    tokens = re.findall(r"[\w]+", normalized, flags=re.UNICODE)
    return [token for token in tokens if len(token) > 1 or token in {"x", "y"}]


def small_talk_answer(question: str) -> Optional[str]:
    normalized = re.sub(r"[^\w\s]", "", question.casefold()).strip()
    normalized = re.sub(r"\s+", " ", normalized)
    greetings = {
        "hi", "hello", "hey", "hai", "hii", "hi there", "hello there", "good morning", "good afternoon", "good evening", "yo",
    }
    thanks = {"thanks", "thank you", "thankyou", "ok", "okay", "cool", "great"}
    identity = {
        "who are you", "what are you", "what can you do", "help", "help me", "what is clario",
    }
    status = {
        "how are you", "how r u", "how are u", "are you there", "status",
    }

    if normalized in greetings:
        return (
            "Hello! I am **Clario**, your document assistant. Ask me anything about your uploaded files, "
            "and I'll synthesize structured, grounded answers with citations."
        )
    if normalized in thanks:
        return "You're welcome! Feel free to ask more questions about your documents."
    if normalized in identity:
        return (
            "I am **Clario**, an AI-powered document research assistant. You can upload PDFs, Word documents, "
            "spreadsheets, text notes, or web articles. I analyze them and answer your questions directly from your data."
        )
    if normalized in status:
        return "All systems operational! My knowledge base and retrieval engine are ready."
    return None


def local_embed(texts: List[str], dimensions: Optional[int] = None) -> List[List[float]]:
    dimensions = dimensions or _embedding_dimension()
    vectors = []
    for text in texts:
        vector = [0.0] * dimensions
        tokens = _tokens(text)
        features = tokens + [f"{a}_{b}" for a, b in zip(tokens, tokens[1:])]

        for feature in features:
            digest = hashlib.blake2b(feature.encode("utf-8"), digest_size=8).digest()
            bucket = int.from_bytes(digest[:4], "big") % dimensions
            sign = 1.0 if digest[4] & 1 else -1.0
            vector[bucket] += sign

        norm = math.sqrt(sum(value * value for value in vector))
        if norm == 0:
            vector[0] = 1.0
        else:
            vector = [value / norm for value in vector]
        vectors.append(vector)
    return vectors


async def embed_texts(texts: List[str], api_key: Optional[str] = None) -> tuple[List[List[float]], str]:
    if not texts:
        return [], "local"

    if is_gemini_configured(api_key):
        try:
            return await gemini_embed(texts, api_key=api_key), "gemini"
        except (GeminiConfigurationError, httpx.HTTPError, OSError) as exc:
            log.warning("Gemini embeddings unavailable; falling back to local embeddings: %s", exc)

    return local_embed(texts), "local"


def store_chunks(chunks: List[Dict], vectors: List[List[float]]) -> int:
    if not chunks:
        return 0
    if len(chunks) != len(vectors):
        raise ValueError("Embedding count did not match chunk count.")

    collection = _get_collection()
    ids = [c["meta"]["id"] for c in chunks]
    documents = [c["text"] for c in chunks]
    metadatas = [c["meta"] for c in chunks]
    collection.add(ids=ids, embeddings=vectors, documents=documents, metadatas=metadatas)
    log.info("Stored %d chunks in ChromaDB", len(chunks))
    return len(chunks)


# ---------------------------------------------------------------------------
# Retrieval & Search Logic
# ---------------------------------------------------------------------------

def lexical_search(question: str, top_k: int = TOP_K_RESULTS) -> List[Dict]:
    collection = _get_collection()
    count = collection.count()
    if count == 0:
        return []

    query_tokens = _tokens(question)
    if not query_tokens:
        return []

    query_counts = Counter(query_tokens)
    query_phrase = " ".join(query_tokens)
    results = collection.get(include=["documents", "metadatas"])
    scored = []
    for doc, meta in zip(results.get("documents", []), results.get("metadatas", [])):
        doc_tokens = _tokens(doc or "")
        if not doc_tokens:
            continue

        doc_counts = Counter(doc_tokens)
        overlap = sum(min(cnt, doc_counts.get(token, 0)) for token, cnt in query_counts.items())
        phrase_bonus = 2 if query_phrase and query_phrase in " ".join(doc_tokens) else 0
        score = (overlap + phrase_bonus) / max(1, len(query_tokens))
        if score <= 0:
            continue

        scored.append({
            "text": doc,
            "source": (meta or {}).get("source", "unknown"),
            "chunk_index": (meta or {}).get("chunk_index", 0),
            "score": round(min(score, 1.0), 4),
        })

    scored.sort(key=lambda hit: hit["score"], reverse=True)
    return scored[:top_k]


def search_similar(query_vec: List[float], top_k: int = TOP_K_RESULTS) -> List[Dict]:
    collection = _get_collection()
    count = collection.count()
    if count == 0:
        return []

    results = collection.query(
        query_embeddings=[query_vec],
        n_results=min(top_k, count),
        include=["documents", "metadatas", "distances"],
    )

    hits = []
    docs = results["documents"][0]
    metas = results["metadatas"][0]
    distances = results["distances"][0]
    for doc, meta, dist in zip(docs, metas, distances):
        score = max(0.0, min(1.0, 1 - dist))
        hits.append({
            "text": doc,
            "source": meta.get("source", "unknown"),
            "chunk_index": meta.get("chunk_index", 0),
            "score": round(score, 4),
        })
    return hits


def merge_hits(*hit_groups: List[Dict], top_k: int = TOP_K_RESULTS) -> List[Dict]:
    merged = {}
    for hits in hit_groups:
        for hit in hits:
            key = (hit.get("source", "unknown"), hit.get("chunk_index", 0))
            if key not in merged or hit.get("score", 0) > merged[key].get("score", 0):
                merged[key] = hit

    ranked = sorted(merged.values(), key=lambda h: h.get("score", 0), reverse=True)
    return ranked[:top_k]


def compact_sources(hits: List[Dict]) -> List[Dict]:
    sources = []
    seen = set()
    for hit in hits:
        source = hit.get("source", "unknown")
        if source in seen:
            continue
        seen.add(source)
        sources.append({"source": source, "score": hit.get("score", 0)})
    return sources


def build_extractive_answer(question: str, hits: List[Dict], reason: Optional[str] = None) -> str:
    if not hits:
        return "I don't have enough information to answer that from the provided documents."

    best = hits[0]
    raw_text = best.get("text") or ""
    note = f"> ℹ️ *{reason}*\n\n" if reason else ""

    # Generate neat bulleted extract of sentences
    sentences = re.split(r"(?<=[.!?])\s+", " ".join(raw_text.split()))
    points = [s.strip(" .") for s in sentences if len(s.strip()) > 25][:5]
    body = "\n".join(f"- {p}" for p in points) if points else raw_text

    return (
        f"{note}"
        f"**Source:** `{best.get('source', 'unknown')}`\n\n"
        f"### Summary of Relevant Findings\n\n"
        f"{body}"
    )


def list_documents() -> List[Dict]:
    collection = _get_collection()
    if collection.count() == 0:
        return []

    results = collection.get(include=["metadatas"])
    counts = Counter(
        meta.get("source", "unknown")
        for meta in results.get("metadatas", [])
        if meta
    )
    return [
        {"source": source, "chunks": chunks}
        for source, chunks in sorted(counts.items(), key=lambda item: item[0].lower())
    ]


def delete_document(source: str) -> int:
    source = source.strip()
    if not source:
        raise ValueError("Document source is required.")

    collection = _get_collection()
    matches = collection.get(where={"source": source}, include=["metadatas"])
    ids = matches.get("ids", [])
    if not ids:
        return 0

    collection.delete(ids=ids)
    log.info("Deleted %d chunks for source: %s", len(ids), source)
    return len(ids)


# ---------------------------------------------------------------------------
# File Extractors & SSRF-Protected Web Scraping
# ---------------------------------------------------------------------------

async def _save_temp_file(upload: UploadFile) -> str:
    suffix = os.path.splitext(upload.filename or "file")[1] or ".bin"
    tmp_path = os.path.join(tempfile.gettempdir(), f"{uuid.uuid4()}{suffix}")
    content = await upload.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds maximum allowed size ({MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB).",
        )
    async with aiofiles.open(tmp_path, "wb") as f:
        await f.write(content)
    return tmp_path


def _extract_pdf(path: str) -> str:
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n".join(pages)


def _extract_docx(path: str) -> str:
    document = docx.Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_csv(path: str) -> str:
    rows = []
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            values = [cell.strip() for cell in row]
            if any(values):
                rows.append(" | ".join(values))
    return "\n".join(rows)


def _validate_safe_url(url: str) -> str:
    """Enforce SSRF prevention by disallowing private/loopback/metadata destinations."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("Enter a valid http or https URL.")

    host = parsed.hostname or ""
    if host.lower() in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("Access to localhost / loopback addresses is not permitted.")

    try:
        addr_info = socket.getaddrinfo(host, None)
        for entry in addr_info:
            ip_str = entry[4][0]
            ip_obj = ipaddress.ip_address(ip_str)
            if (
                ip_obj.is_private
                or ip_obj.is_loopback
                or ip_obj.is_reserved
                or ip_obj.is_link_local
            ):
                raise ValueError(f"Access to private/internal network address ({ip_str}) is blocked.")
    except socket.gaierror:
        raise ValueError(f"Unable to resolve hostname: {host}")

    return url.strip()


async def _scrape_url(url: str) -> str:
    current_url = _validate_safe_url(url)
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Clario/2.1"}

    # Manually follow redirects up to 5 hops, re-validating the destination IP on every hop
    async with httpx.AsyncClient(timeout=15, follow_redirects=False) as client:
        resp = None
        for _ in range(5):
            resp = await client.get(current_url, headers=headers)
            if resp.is_redirect:
                redirect_target = resp.headers.get("Location")
                if not redirect_target:
                    break
                from urllib.parse import urljoin
                current_url = _validate_safe_url(urljoin(current_url, redirect_target))
                continue
            resp.raise_for_status()
            break
        else:
            raise ValueError("Too many redirects encountered while scraping URL.")

    if resp is None:
        raise ValueError("Could not complete request to URL.")

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "form", "nav", "footer", "header", "aside"]):
        tag.decompose()

    blocks = []
    seen = set()
    for tag in soup.find_all(["h1", "h2", "h3", "p", "li"]):
        text = " ".join(tag.get_text(separator=" ", strip=True).split())
        if len(text) < 25 or text in seen:
            continue
        seen.add(text)
        blocks.append(text)

    return "\n\n".join(blocks)


async def process_upload(file: UploadFile, api_key: Optional[str] = None) -> int:
    tmp_path = await _save_temp_file(file)
    suffix = os.path.splitext(tmp_path)[1].lower()

    try:
        # Offload CPU-bound parsing to worker thread
        if suffix == ".pdf":
            raw = await asyncio.to_thread(_extract_pdf, tmp_path)
        elif suffix == ".docx":
            raw = await asyncio.to_thread(_extract_docx, tmp_path)
        elif suffix in {".txt", ".md"}:
            raw = await asyncio.to_thread(_extract_txt, tmp_path)
        elif suffix == ".csv":
            raw = await asyncio.to_thread(_extract_csv, tmp_path)
        else:
            raise ValueError(f"Unsupported file type: '{suffix}'. Supported: PDF, DOCX, TXT, MD, CSV")
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    if not raw.strip():
        raise ValueError("No text could be extracted from the file.")

    chunks = split_text(raw, source_file=file.filename or "upload")
    vectors, provider = await embed_texts([c["text"] for c in chunks], api_key=api_key)
    for chunk in chunks:
        chunk["meta"]["embedding_provider"] = provider
    return store_chunks(chunks, vectors)


async def process_url(url: str, api_key: Optional[str] = None) -> int:
    raw = await _scrape_url(url)
    if not raw.strip():
        raise ValueError("No text could be extracted from the URL.")

    chunks = split_text(raw, source_file=url)
    vectors, provider = await embed_texts([c["text"] for c in chunks], api_key=api_key)
    for chunk in chunks:
        chunk["meta"]["embedding_provider"] = provider
    return store_chunks(chunks, vectors)


# ---------------------------------------------------------------------------
# API Routes
# ---------------------------------------------------------------------------

@router.get("/health")
async def health(x_gemini_key: Optional[str] = Header(None, alias="X-Gemini-Key")):
    collection = _get_collection()
    has_gemini = is_gemini_configured(x_gemini_key)
    return {
        "status": "ok",
        "service": "Clario API",
        "version": "2.1.0",
        "model": LLM_MODEL,
        "gemini_configured": has_gemini,
        "vector_db": "ChromaDB (local persistent)",
        "chunks_stored": collection.count(),
        "documents_count": len(list_documents()),
    }


@router.get("/documents", response_model=DocumentsResponse)
async def documents():
    docs = list_documents()
    return DocumentsResponse(
        documents=[DocumentSummary(**doc) for doc in docs],
        total_documents=len(docs),
        total_chunks=sum(doc["chunks"] for doc in docs),
    )


@router.delete("/documents", response_model=DeleteDocumentResponse)
async def remove_document(body: DeleteDocumentRequest):
    try:
        deleted = delete_document(body.source)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Document was not found in the knowledge base.")

        return DeleteDocumentResponse(
            source=body.source,
            chunks_deleted=deleted,
            documents_remaining=len(list_documents()),
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except HTTPException:
        raise
    except Exception as exc:
        log.exception("Document deletion failed")
        raise HTTPException(status_code=500, detail=f"Document deletion error: {exc}")


@router.post("/ingest", response_model=IngestResponse)
async def ingest_file(
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    x_gemini_key: Optional[str] = Header(None, alias="X-Gemini-Key"),
):
    try:
        upload_files = files or ([file] if file is not None else [])
        if not upload_files:
            raise ValueError("No files were uploaded.")

        total_chunks = 0
        filenames = []
        for upload in upload_files:
            log.info("Ingesting file: %s", upload.filename)
            total_chunks += await process_upload(upload, api_key=x_gemini_key)
            filenames.append(upload.filename or "upload")

        return IngestResponse(
            status="success",
            chunks_stored=total_chunks,
            filename=", ".join(filenames),
            filenames=filenames,
            files_processed=len(upload_files),
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except HTTPException:
        raise
    except GeminiConfigurationError as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"Embedding provider request failed: {exc}")
    except Exception as exc:
        log.exception("Ingestion failed")
        raise HTTPException(status_code=500, detail=f"Ingestion error: {exc}")


@router.post("/ingest-url", response_model=IngestResponse)
async def ingest_url(
    body: IngestURLRequest,
    x_gemini_key: Optional[str] = Header(None, alias="X-Gemini-Key"),
):
    try:
        log.info("Scraping URL: %s", body.url)
        count = await process_url(body.url, api_key=x_gemini_key)
        return IngestResponse(
            status="success",
            chunks_stored=count,
            filename=body.url,
            filenames=[body.url],
            files_processed=1,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except HTTPException:
        raise
    except GeminiConfigurationError as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    except httpx.HTTPStatusError as exc:
        raise HTTPException(status_code=502, detail=f"URL returned HTTP {exc.response.status_code}.")
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"URL scrape request failed: {exc}")
    except Exception as exc:
        log.exception("URL ingestion failed")
        raise HTTPException(status_code=500, detail=f"URL scrape error: {exc}")


@router.post("/chat", response_model=ChatResponse)
async def chat(
    payload: ChatRequest,
    x_gemini_key: Optional[str] = Header(None, alias="X-Gemini-Key"),
):
    if not payload.question.strip():
        raise HTTPException(status_code=400, detail="Question cannot be empty.")

    try:
        canned = small_talk_answer(payload.question)
        if canned:
            return ChatResponse(answer=canned, sources=[], chunks=[])

        if _get_collection().count() == 0:
            return ChatResponse(
                answer="No documents are currently indexed in Clario. Please upload a PDF, DOCX, CSV, or web link to start chatting!",
                sources=[],
                chunks=[],
            )

        lexical_hits = lexical_search(payload.question, top_k=TOP_K_RESULTS)
        query_vectors, query_provider = await embed_texts([payload.question], api_key=x_gemini_key)
        vector_hits = search_similar(query_vectors[0], top_k=TOP_K_RESULTS)
        hits = merge_hits(lexical_hits, vector_hits, top_k=TOP_K_RESULTS)

        if not hits:
            return ChatResponse(
                answer="I don't have enough information to answer that from the provided documents.",
                sources=[],
                chunks=[],
            )

        if query_provider != "gemini":
            fallback_hits = lexical_hits or hits
            return ChatResponse(
                answer=build_extractive_answer(
                    payload.question,
                    fallback_hits,
                    "Gemini is offline or unconfigured. Below is the relevant extracted passage.",
                ),
                sources=compact_sources(fallback_hits),
                chunks=fallback_hits,
            )

        full_prompt = build_prompt(
            retrieved_chunks=hits,
            chat_history=payload.chat_history,
            question=payload.question,
        )

        try:
            answer, _ = await gemini_chat(full_prompt, model=LLM_MODEL, api_key=x_gemini_key)
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code == 429:
                answer = build_extractive_answer(
                    payload.question,
                    hits,
                    "Gemini API is currently rate-limited (429). Here is the relevant extracted passage.",
                )
            else:
                raise
        except (httpx.HTTPError, OSError) as exc:
            answer = build_extractive_answer(
                payload.question,
                hits,
                f"Gemini API request failed: {exc}",
            )

        return ChatResponse(
            answer=answer,
            sources=compact_sources(hits),
            chunks=hits,
        )
    except HTTPException:
        raise
    except GeminiConfigurationError as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    except Exception as exc:
        log.exception("Chat failed")
        raise HTTPException(status_code=500, detail=f"Chat error: {exc}")


@router.post("/chat/stream")
async def chat_stream(
    payload: ChatRequest,
    x_gemini_key: Optional[str] = Header(None, alias="X-Gemini-Key"),
):
    """
    Server-Sent Events (SSE) streaming chat endpoint.
    Emits real-time tokens directly to the browser for zero perceived latency.
    """
    if not payload.question.strip():
        raise HTTPException(status_code=400, detail="Question cannot be empty.")

    async def event_generator() -> AsyncGenerator[str, None]:
        canned = small_talk_answer(payload.question)
        if canned:
            yield f"data: {json.dumps({'type': 'sources', 'sources': []})}\n\n"
            yield f"data: {json.dumps({'type': 'delta', 'text': canned})}\n\n"
            yield f"data: {json.dumps({'type': 'done'})}\n\n"
            return

        if _get_collection().count() == 0:
            yield f"data: {json.dumps({'type': 'sources', 'sources': []})}\n\n"
            yield f"data: {json.dumps({'type': 'delta', 'text': 'No documents are currently indexed in Clario. Please upload documents in the sidebar to begin!'})}\n\n"
            yield f"data: {json.dumps({'type': 'done'})}\n\n"
            return

        lexical_hits = lexical_search(payload.question, top_k=TOP_K_RESULTS)
        query_vectors, query_provider = await embed_texts([payload.question], api_key=x_gemini_key)
        vector_hits = search_similar(query_vectors[0], top_k=TOP_K_RESULTS)
        hits = merge_hits(lexical_hits, vector_hits, top_k=TOP_K_RESULTS)

        sources = compact_sources(hits)
        yield f"data: {json.dumps({'type': 'sources', 'sources': sources})}\n\n"

        if not hits:
            no_info_msg = "I don't have enough information to answer that from the provided documents."
            yield f"data: {json.dumps({'type': 'delta', 'text': no_info_msg})}\n\n"
            yield f"data: {json.dumps({'type': 'done'})}\n\n"
            return

        # Fallback if Gemini not available
        if query_provider != "gemini":
            fallback_text = build_extractive_answer(
                payload.question,
                hits,
                "Gemini is offline or unconfigured. Below is the relevant extracted passage.",
            )
            # Stream in friendly word chunks
            for word in fallback_text.split(" "):
                yield f"data: {json.dumps({'type': 'delta', 'text': word + ' '})}\n\n"
                await asyncio.sleep(0.015)
            yield f"data: {json.dumps({'type': 'done'})}\n\n"
            return

        full_prompt = build_prompt(
            retrieved_chunks=hits,
            chat_history=payload.chat_history,
            question=payload.question,
        )

        try:
            async for token in gemini_chat_stream(full_prompt, model=LLM_MODEL, api_key=x_gemini_key):
                yield f"data: {json.dumps({'type': 'delta', 'text': token})}\n\n"
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code == 429:
                fallback_msg = build_extractive_answer(
                    payload.question,
                    hits,
                    "Gemini API rate limit reached (429). Here is the relevant extracted text.",
                )
                yield f"data: {json.dumps({'type': 'delta', 'text': fallback_msg})}\n\n"
            else:
                yield f"data: {json.dumps({'type': 'delta', 'text': f'Gemini error: {exc}'})}\n\n"
        except Exception as exc:
            yield f"data: {json.dumps({'type': 'delta', 'text': f'Generation error: {exc}'})}\n\n"

        yield f"data: {json.dumps({'type': 'done'})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


# Attach API router
app.include_router(router, prefix="/api")


# ---------------------------------------------------------------------------
# Static Web App Mount (Production Unified Serving)
# ---------------------------------------------------------------------------

dist_candidates = [
    os.path.join(BASE_DIR, "..", "frontend", "dist"),
    os.path.join(BASE_DIR, "frontend", "dist"),
    os.path.join(BASE_DIR, "dist"),
]
dist_dir = next((path for path in dist_candidates if os.path.isdir(path)), None)

if dist_dir:
    log.info("Mounting built frontend static files from: %s", dist_dir)
    assets_dir = os.path.join(dist_dir, "assets")
    if os.path.isdir(assets_dir):
        app.mount("/assets", StaticFiles(directory=assets_dir), name="assets")

    @app.get("/{full_path:path}")
    async def serve_spa(full_path: str):
        if full_path.startswith("api/") or full_path == "docs" or full_path == "openapi.json":
            raise HTTPException(status_code=404)
        target = os.path.join(dist_dir, full_path)
        if os.path.isfile(target):
            return FileResponse(target)
        return FileResponse(os.path.join(dist_dir, "index.html"))


if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    uvicorn.run("rag_chatbot:app", host="0.0.0.0", port=port, reload=True)
